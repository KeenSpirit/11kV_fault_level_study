import powerfactory as pf
import tkinter as tk
from tkinter import ttk
import math
import sys
from importlib import reload
import user_inputs as ui
import floating_terminals
import lines_results
import fdr_open_points
reload(lines_results)
reload(ui)
reload(fdr_open_points)
reload(floating_terminals)

def main(app):
    """
    11kV fault level study
    """

    #
    # INITIALISE SCRIPT
    #

    # Enables the user to manually stop the script
    app.SetEnableUserBreak(1)

    # Store the initial feeder switch state
    switch_state = {}
    stored_switch_state = store_switch_state(app, switch_state, 0)

    # Model correction
    # oos_lines(app)

    #
    # DATA COLLECTION
    #

    sub_name = substation_name(app)
    feeder_list, external_grid = ui.get_input(app)
    feeders_switches = get_feeder_switches(app, feeder_list)
    feeders_relays = run_window(app, feeder_list, feeders_switches)
    feeders_devices_terminals_objects = get_downstream_devices(app, feeders_relays)
    _, feeders_devices_load_total = section_capacity(app, feeders_devices_terminals_objects)
    feeders_devices_section = get_device_sections(app, feeders_devices_terminals_objects)
    feeders_sections_loads, _ = section_capacity(app, feeders_devices_section)
    feeders_sections_loading, feeders_sections_trmax, feeders_devices_loads = tr_data(app, feeders_sections_loads,
                                                                                      feeders_devices_section)
    feeders_sections_load_total = reorder_dictionary_1(feeders_devices_section, feeders_devices_load_total)
    floating_terms = get_floating_terminals(app, feeder_list, feeders_devices_section)

    #
    # CALCULATE RESULTS
    #

    # Switch out unstudied feeders
    # switch_feeders(app, radial_list, feeder_list)

    # Inrush
    feeders_devices_inrush, = inrush(feeders_sections_load_total)

    app.PrintPlain("Calculating max fault levels...")

    short_circuit(app, 3, 'spgf', 0)
    # max tfmr PG fault level
    max_tr_pg = target_values(feeders_sections_trmax, 1)
    max_tr_pg_term, max_tr_pg_object = result_bounds(max_tr_pg, 1)
    # phase-ground
    first_pass_max_pg = target_values(feeders_devices_section, 1)
    unordered_all_max_pg, results_all_max_pg_obj = append_floating_terms(app, first_pass_max_pg, floating_terms)
    results_max_pg, _ = result_bounds(results_all_max_pg_obj, 1)

    # 2 phase
    short_circuit(app, 3, '2psc', 0)
    first_pass_max_2p = target_values(feeders_devices_section, 2)
    unordered_all_max_2p, unordered_all_max_2p_obj = append_floating_terms(app, first_pass_max_2p, floating_terms)
    results_max_2p, _ = result_bounds(unordered_all_max_2p_obj, 1)

    short_circuit(app, 3, '3psc', 0)
    # max tfmr 3P fault level
    max_tr_3p = target_values(feeders_sections_trmax, 3)
    max_tr_3p_term, max_tr_3p_object = result_bounds(max_tr_3p, 1)
    # 3 phase
    first_pass_max_3p = target_values(feeders_devices_section, 3)
    unordered_all_max_3p, unordered_all_max_3p_obj = append_floating_terms(app, first_pass_max_3p, floating_terms)
    results_max_3p, _ = result_bounds(unordered_all_max_3p_obj, 1)

    app.PrintPlain("Calculating min fault levels...")
    # 2 phase
    short_circuit(app, 3, '2psc', 1)
    first_pass_min_2p = target_values(feeders_devices_section, 2)
    unordered_all_min_2p, unordered_all_min_2p_obj = append_floating_terms(app, first_pass_min_2p, floating_terms)
    results_min_2p, _ = result_bounds(unordered_all_min_2p_obj, 0)

    # 3 phase
    short_circuit(app, 3, '3psc', 1)
    first_pass_min_3p = target_values(feeders_devices_section, 3)
    unordered_all_min_3p, unordered_all_min_3p_obj = append_floating_terms(app, first_pass_min_3p, floating_terms)
    results_min_3p, _ = result_bounds(unordered_all_min_3p_obj, 0)

    # phase-ground
    short_circuit(app, 3, 'spgf', 1)
    first_pass_min_pg = target_values(feeders_devices_section, 1)
    results_all_min_pg, unordered_all_min_pg_obj = append_floating_terms(app, first_pass_min_pg, floating_terms)
    results_min_pg, _ = result_bounds(unordered_all_min_pg_obj, 0)

    app.PrintPlain("Calculating system normal min fault levels...")
    reset_min_source_imp(external_grid, sys_norm_min=True)
    # 2 phase
    short_circuit(app, 3, '2psc', 1)
    first_pass_min_2p = target_values(feeders_devices_section, 2)
    unordered_all_sys_norm_min_2p, unordered_all_min_2p_obj = append_floating_terms(app, first_pass_min_2p, floating_terms)
    result_sys_norm_min_2p, _ = result_bounds(unordered_all_min_2p_obj, 0)

    # phase-ground
    short_circuit(app, 3, 'spgf', 1)
    first_pass_min_pg = target_values(feeders_devices_section, 1)
    unordered_all_sys_norm_min_pg, unordered_all_min_pg_obj = append_floating_terms(app, first_pass_min_pg, floating_terms)
    result_sys_norm_min_pg, _ = result_bounds(unordered_all_min_pg_obj, 0)

    # Re-order dictionaries to align with min pg fault levels
    results_all_max_2p = reorder_dictionary_2(results_all_min_pg, unordered_all_max_2p)
    results_all_max_3p = reorder_dictionary_2(results_all_min_pg, unordered_all_max_3p)
    results_all_max_pg = reorder_dictionary_2(results_all_min_pg, unordered_all_max_pg)
    results_all_min_2p = reorder_dictionary_2(results_all_min_pg, unordered_all_min_2p)
    results_all_min_3p = reorder_dictionary_2(results_all_min_pg, unordered_all_min_3p)
    result_all_sys_norm_min_2p = reorder_dictionary_2(results_all_min_pg, unordered_all_sys_norm_min_2p)
    result_all_sys_norm_min_pg = reorder_dictionary_2(results_all_min_pg, unordered_all_sys_norm_min_pg)

    reset_min_source_imp(external_grid)

    # Transformer data
    results_max_tr_3p, feeders_sections_trmax_size = largest_tr(app, feeders_sections_loading, max_tr_3p_object)
    results_max_tr_pg, _ = largest_tr(app, feeders_sections_loading, max_tr_pg_object)

    feeders_devices_load = reorder_dictionary_2(results_all_max_3p, feeders_devices_loads)

    # Open point data
    fdrs_open_switches = {}
    fdrs_lines_oos = {}
    for feeder in feeder_list:
        open_switches, lines_oos = fdr_open_points.get_open_points(app, feeder)
        fdrs_open_switches[feeder] = open_switches
        fdrs_lines_oos[feeder] = lines_oos

    # Obtain line results for conductor damage studies
    results_lines_max_3p, results_lines_max_2p, results_lines_max_pg, results_lines_min_3p, results_lines_min_2p, \
    results_lines_min_pg, result_lines_sys_norm_min_2p, result_lines_sys_norm_min_pg, result_lines_type, \
    result_lines_therm_rating = lines_results.lines(app, results_all_max_3p, results_all_max_2p, results_all_max_pg,
    results_all_min_3p, results_all_min_2p, results_all_min_pg, result_all_sys_norm_min_2p, result_all_sys_norm_min_pg)

    #
    # PASS RESULTS TO OUTPUT FILE
    #
    ext_grid_convert = {}
    for grid, value in external_grid.items():
        ext_grid_convert[grid.GetAttribute('loc_name')] = value


    output = output_results(app, sub_name, ext_grid_convert, feeders_devices_inrush,
            results_max_3p, results_max_2p, results_max_pg, results_min_2p, results_min_3p, results_min_pg,
            result_sys_norm_min_2p, result_sys_norm_min_pg, feeders_sections_trmax_size, results_max_tr_3p,
            results_max_tr_pg, results_all_max_3p, results_all_max_2p, results_all_max_pg, results_all_min_3p,
            results_all_min_2p, results_all_min_pg, result_all_sys_norm_min_2p, result_all_sys_norm_min_pg,
            feeders_devices_load, results_lines_max_3p, results_lines_max_2p, results_lines_max_pg,
            results_lines_min_3p, results_lines_min_2p, results_lines_min_pg, result_lines_sys_norm_min_2p,
            result_lines_sys_norm_min_pg, result_lines_type, result_lines_therm_rating, fdrs_open_switches, fdrs_lines_oos)

    save_results(app, sub_name, output)

    #
    # HOUSEKEEPING
    #

    # Restore feeder switches to the initial state
    store_switch_state(app, stored_switch_state, 1)

    # script run successfully.


def store_switch_state(app, switch_state, state):
    """Store or restore the initial model switch state. If state = 0, the initial switch state is stored. If state != 0,
    the initial switch state is restored"""

    all_switches = app.GetCalcRelevantObjects('*.ElmCoup')

    if state == 0:
        # Store the initial feeder switch state
        for switch in all_switches:
            switch_state[switch] = switch.on_off
    else:
        # Restore feeder switches to the initial state before ending script
        for switch in all_switches:
            switch.on_off = switch_state[switch]

    return switch_state


# def oos_lines(app):
#     """
#     From a list of lines, find any lines that are out of service. Place these lines in service and place a switch on
#     one end of the line. Isolate the line by opening the switch.
#     """
#     all_oos_lines = []
#     for grid in app.GetSummaryGrid().GetContents():
#         all_oos_lines += [
#             line
#             for line in grid.obj_id.GetContents("*.ElmLne")
#             if "HV" in line.loc_name or "TR" in line.loc_name
#                or "LN" in line.loc_name
#             if line.IsOutOfService()
#         ]
#     for line in all_oos_lines:
#         if line.outserv == 1:
#             line.outserv = 0
#             switches_1 = line.bus1.GetContents("*.StaSwitch")
#             switches_2 = line.bus2.GetContents("*.StaSwitch")
#             if switches_1:
#                 for switch in switches_1:
#                     switch.on_off = 0
#             elif switches_2:
#                 for switch in switches_2:
#                     switch.on_off = 0
#             elif line.bus2:
#                 new_switch = line.bus2.CreateObject("StaSwitch", "temp_switch")
#                 try:
#                     new_switch.on_off = 0
#                 except:
#                     app.PrintPlain(line)
#             else:
#                 new_switch = line.bus1.CreateObject("StaSwitch", "temp_switch")
#                 new_switch.on_off = 0


def substation_name(app):
    """Get the substation name"""

    sub_names = app.GetCalcRelevantObjects('*.ElmNet')
    subs = []
    for sub in sub_names:
        if sub.loc_name != "New Elements" and sub.loc_name != "Summary Grid" and sub.loc_name != "Boundary Subs":
            subs.append(sub.loc_name + " 11kV")
    if len(subs) > 20:
        app.PrintPlain("Too many substation grids are active. "
                       "Please ensure twenty or less grids are active and rerun the script.")
        sys.exit(0)
    else:
        sub_name = '_'.join(subs)
    return sub_name


def get_feeder_switches(app, feeder_list):
    """Get a lists of all switches for all feeders"""

    all_switches = app.GetCalcRelevantObjects("*.StaSwitch")

    feeders_switches = {}
    for feeder in feeder_list:
        cache = []
        feeder_elm = app.GetCalcRelevantObjects(feeder + ".ElmFeeder")[0]
        terminal_list = feeder_elm.GetObjs('ElmTerm')
        for switch in all_switches:
            # Get the terminal associated with the StaSwitch.
            cubicle = switch.fold_id
            switch_terminal = cubicle.cterm
            # If that terminal is in the feeder terminal_list list, add it to the feeder switch list.
            if switch.on_off == 1 and switch_terminal in terminal_list:
                cache.append(switch)

        # Sort the list of switches
        def name(switch_obj):
            return switch_obj.loc_name

        cache.sort(key=name)
        feeders_switches[feeder] = cache

    return feeders_switches


def populate(root, frame, feeder_list, feeders_switches):
    """Put data for protection device selection into frame"""

    # If a switch has a "_Term" suffix, chop it off
    feeders_switches_chop = {feeder: [] for feeder in feeders_switches}
    for feeder, switches in feeders_switches.items():
        for switch in switches:
            cubicle = switch.fold_id
            switch_term = cubicle.cterm.loc_name
            if switch_term[-5:] == "_Term":
                feeders_switches_chop[feeder].append(switch_term[:-5])
            else:
                feeders_switches_chop[feeder].append(switch_term)

    ttk.Label(frame, text="Identify ALL PMRs and fuses belonging to the given feeder:", font='Helvetica 12 bold').grid(
        columnspan=8, padx=5, pady=5)
    for fid in feeder_list:
        ttk.Label(frame, text=fid).grid(row=1, column=feeder_list.index(fid), sticky='W', padx=10, pady=5)

    # Create the checkbox

    var = []
    max_list_length = 0
    for feeder in feeders_switches_chop:
        list_length = len(feeders_switches_chop[feeder])
        if list_length > max_list_length:
            max_list_length = list_length
        switch_list = feeders_switches_chop[feeder]
        for i in range(list_length):
            var.append(tk.IntVar())
            ttk.Checkbutton(frame, text=switch_list[i], variable=var[-1]).grid(row=i + 4, column=list(feeders_switches_chop).index(feeder), sticky='W', padx=10, pady=5)
    ttk.Button(frame, text='Okay', command=lambda: root.destroy()).grid(row=max_list_length + 5, column=0, sticky='W',padx=5, pady=5)
    ttk.Button(frame, text='Exit', command=lambda: exit_script(root)).grid(row=max_list_length + 5, column=1, sticky='W', pady=5)

    return var, feeders_switches_chop


def onFrameConfigure(canvas):
    """Reset the scroll region to encompass the inner frame"""
    canvas.configure(scrollregion=canvas.bbox("all"))


def run_window(app, feeder_list, feeders_switches):
    """Build the canvas and embed the frame"""

    def _window_dim(feeder_list, feeders_switches):
        horiz_offset = 500
        num_columns = len(feeder_list)
        column_width = 120
        col_padding = 0
        window_width = max((num_columns * column_width + col_padding), 550)
        if window_width > 1300:
            horiz_offset = 200
            window_width = 1500
        num_rows = max(len(list) for list in feeders_switches.values())
        row_height = 32
        row_padding = 100
        window_height = num_rows * row_height + row_padding
        if window_height > 900:
            window_height = 900
        return window_width, window_height, horiz_offset

    root = tk.Tk()
    root.title("11kV Fault Study")
    window_width, window_height, horiz_offset = _window_dim(feeder_list, feeders_switches)
    root.geometry(f"{window_width}x{window_height}+{horiz_offset}+100")

    canvas = tk.Canvas(root, borderwidth=0)
    frame = tk.Frame(canvas)
    vsb = tk.Scrollbar(root, orient="vertical", command=canvas.yview)
    canvas.configure(yscrollcommand=vsb.set)

    vsb.pack(side="right", fill="y")
    canvas.pack(side="left", fill="both", expand=True)
    canvas.create_window((4, 4), window=frame, anchor="nw")

    frame.bind("<Configure>", lambda event, canvas=canvas: onFrameConfigure(canvas))

    var, feeders_switches_chop = populate(root, frame, feeder_list, feeders_switches)

    root.mainloop()

    # Get the list of PMRs and fuses
    pmr_fuse_list = []
    count = 0
    for feeder in feeders_switches_chop:
        bog = len(feeders_switches_chop[feeder])
        switch_list = feeders_switches[feeder]
        for i in range(bog):
            if var[count].get() == 1:
                pmr_fuse_list.append(switch_list[i])
            count += 1

    # this comprehension removes switches if not PMRs or fuses. It does membership checks of pmr_fuse_list by iterating
    # over members of feeders_switches
    feeders_relays = {i: [a for a in j if a in pmr_fuse_list] for i, j in feeders_switches.items()}

    # append the feeder name to the PMR list
    for feeder in feeders_relays.keys():
        feeder_elm = app.GetCalcRelevantObjects(feeder + ".ElmFeeder")
        feeders_relays[feeder].append(feeder_elm[0])

    return feeders_relays


def exit_script(root):
    """Exits script"""

    app = pf.GetApplication()
    app.PrintPlain("User terminated script.")
    root.destroy()
    sys.exit(0)


def get_downstream_devices(app, feeders_devices):
    """For a list of devices within a list of feeders, determine the downstream devices"""

    app.PrintPlain("Building section lists...")
    all_grids = app.GetCalcRelevantObjects('*.ElmXnet')
    grids = [grid for grid in all_grids if grid.outserv == 0]
    all_feeders = app.GetCalcRelevantObjects("*.ElmFeeder")

    feeders_devices_terminals_objects = {feeder: {} for feeder in feeders_devices}
    for feeder, devices in feeders_devices.items():
        for device in devices:
            # Convert device to its equivalent terminal element
            if device.GetClassName() == "ElmFeeder":
                cubicle = device.obj_id
            else:
                cubicle = device.fold_id        # StaSwitch
            device_01 = cubicle.cterm
            # Get the loc_name of the terminal and chop
            # The way you find the terminal associated with Elm.Feeder is different to that
            # for StaSwitch (see NRG2A for example).
            if device in all_feeders:
                device_03 = device.loc_name
            else:
                device_02 = device_01.loc_name
                if device_02[-5:] == "_Term":
                    device_03 = device_02[:-5]
                elif device_02[-8:] == "_Term(1)":
                    device_03 = device_02[:-8]
                else:
                    device_03 = device_02
            # Do a topological search of the device downstream devices
            down_devices = cubicle.GetAll(1, 0)
            # If the external grid is in the downstream list, you're searching in the wrong direction
            if any(item in grids for item in down_devices):
                device_list = cubicle.GetAll(0, 0)
            else:
                device_list = down_devices
            # Filter for terminals
            downstream_list_objects = []
            for device_k in device_list:
                if (device_k.GetClassName() == "ElmTerm"
                        and device_k.loc_name != feeder + "_Term"
                        and device_k.uknom == 11.):
                    downstream_list_objects.append(device_k)
            if not downstream_list_objects:
                app.PrintPlain("The script cannot determine the protection section of " + device.loc_name +
                               " because the circuit path is looped.")
                app.PrintPlain("Please radialise the device or remove it from the study list.")
                sys.exit(0)
            else:
                feeders_devices_terminals_objects[feeder][device_03] = downstream_list_objects

    # append the device name to the device terminal list
    for feeder, devices in feeders_devices_terminals_objects.items():
        for device in devices:
            device_object = app.GetCalcRelevantObjects(device+"*.ElmTerm")
            try:
                device_obj = device_object[0]
            except IndexError:
                device_elm = app.GetCalcRelevantObjects(device + ".ElmFeeder")
                device_cubicle = device_elm[0].obj_id
                device_obj = device_cubicle.cterm
            feeders_devices_terminals_objects[feeder][device].append(device_obj)

    return feeders_devices_terminals_objects


def section_capacity(app, feeders_devices_terminals_objects):
    """For a dictionary of devices with downstream terminals, get downstream load capacity"""

    # For every load in the study case, store its terminal and its rating as a dictionary
    all_loads = app.GetCalcRelevantObjects('*.ElmLod')
    terminals_loads = {}
    for load in all_loads:
        try:
            if load.iLoadTrf == 1:
                rating = round(load.Strat)
                load_cubicle = load.bus1
                term = load_cubicle.cterm
                terminals_loads[term] = rating
        except AttributeError:
            continue

    # For each item in the list of terminals, determine if it's a load by checking it against the above dictionary
    # If it is a load, add the load to a counter of the total device downstream load
    feeders_devices_loads = {feeder: {device: [] for device in feeders_devices_terminals_objects[feeder]} for feeder in
                            feeders_devices_terminals_objects}
    feeders_devices_load_total = {feeder: {} for feeder in feeders_devices_terminals_objects}
    for feeder, devices in feeders_devices_terminals_objects.items():
        for device, terminals in devices.items():
            total_load = 0
            for terminal in terminals:
                if terminals_loads.get(terminal) is not None:
                    feeders_devices_loads[feeder][device].append(terminal)
                    total_load += terminals_loads[terminal]
            feeders_devices_load_total[feeder][device] = total_load

    return feeders_devices_loads, feeders_devices_load_total


def get_device_sections(app, feeders_devices_terminals: dict) -> dict:
    """For a list of devices within a list of feeders, determine the device sections"""

    feeders_devices_sections = {feeder: {} for feeder in feeders_devices_terminals}
    for feeder, devices in feeders_devices_terminals.items():
        short_iter = []
        # while there are still device lists in the feeder key...
        while len(devices) > 0:
            # find the length of the shortest list of terminals for the given feeder.
            min_val = min([len(devices[ele]) for ele in devices])
            # find the shortest list by comparing list lengths to the length of the shortest list.
            for ele in devices:
                if len(devices[ele]) == min_val:
                    shortest_list = ele  # this is the key for the shortest list (ie name of PMR)
            # pop the shortest list dictionary from the dictionary
            short = devices.pop(shortest_list)
            shortest_section = list(set(short) - set(short_iter))

            # chop the device value from its section list so upstream section will include it in its list
            for n in range(min_val):
                try:
                    shortest_list_object = app.GetCalcRelevantObjects(shortest_list + "_Term.ElmTerm")
                    if short[n] == shortest_list_object[0]:
                        del short[n]
                except IndexError:
                    continue

            short_iter = short_iter + short
            feeders_devices_sections[feeder][shortest_list] = shortest_section

    return feeders_devices_sections


def tr_data(app, feeders_sections_loads, feeders_devices_section):
    """From a list of devices within a feeder section, return the capacity of every transformer and
    the names of the largest trs"""

    loads = app.GetCalcRelevantObjects('*.ElmLod')
    # Give capacity of each downstream transformer
    feeders_sections_loading = {feeder: {section: {} for section in feeders_sections_loads[feeder]} for feeder in
                            feeders_sections_loads}
    for feeder, sections in feeders_sections_loads.items():
        for section, stations in sections.items():
            for station in stations:
                for load in loads:
                    try:
                        termination = load.bus1.cterm               # Match the ElmLod with its corresponding ElmTerm
                    except AttributeError:
                        continue
                    if termination == station:                      # by linking them through their cubicle
                        rated_power = load.Strat
                        feeders_sections_loading[feeder][section][station] = round(rated_power)

    feeders_devices_loads = {feeder: {} for feeder in feeders_devices_section}
    for feeder, sections in feeders_devices_section.items():
        feeders_devices_loads[feeder] = {}
        for section, stations in sections.items():
            feeders_devices_loads[feeder][section] = {}
            for station in stations:
                if station in feeders_sections_loading[feeder][section]:
                    feeders_devices_loads[feeder][section][station.loc_name] = feeders_sections_loading[feeder][section][station]
                else:
                    feeders_devices_loads[feeder][section][station.loc_name] = ""

    # For each section, give the terminal with the largest transformer
    feeders_sections_trmax = {feeder: {section: [] for section in feeders_sections_loading[feeder]} for feeder in
                            feeders_sections_loading}
    for feeder, sections in feeders_sections_loading.items():
        for section, trs in sections.items():
            all_values = trs.values()
            try:
                max_value = max(all_values)
            except ValueError:
                max_value = 0.1
            for key, value in trs.items():
                if value == max_value:
                    feeders_sections_trmax[feeder][section].append(key)

    # For the largest transformer, find the terminal in its parent folder corresponding to the original terminals list.
    # This is the terminal that will have short circuit results.
    sec_subs = app.GetCalcRelevantObjects('*.ElmTrfstat')
    feeders_sections_trmax_allterms = {feeder: {device: [] for device in feeders_sections_trmax[feeder]} for feeder in
                             feeders_sections_trmax}
    for feeder, devices in feeders_sections_trmax.items():
        for device, terminals in devices.items():
            for sec_sub in sec_subs:
                term_contents = sec_sub.GetContents('*.ElmTerm', 1)
                if any(item in term_contents for item in terminals) is True:
                    for term in term_contents:
                        if term in feeders_devices_section[feeder][device]:
                            feeders_sections_trmax_allterms[feeder][device].append(term)

    # Sometimes the parent folder won't include an ElmTerm object (rare). In this case, just use the given terminal.
    for feeder, sections in feeders_sections_trmax_allterms.items():
        for section in sections:
            if not feeders_sections_trmax_allterms[feeder][section]:
                feeders_sections_trmax_allterms[feeder][section] = feeders_sections_trmax[feeder][section]

    return feeders_sections_loading, feeders_sections_trmax_allterms, feeders_devices_loads


def reorder_dictionary_1(reference_dictionary, dictionary):
    """Re-order a nested dictionary for output formatting in the Results Summary tab"""

    reordered_dictionary = {feeder: {} for feeder in reference_dictionary}
    for feeder, sections in reference_dictionary.items():
        for section in sections:
            reordered_dictionary[feeder][section] = dictionary[feeder][section]

    return reordered_dictionary


def reorder_dictionary_2(reference_dictionary, dictionary):
    """Re-order a second order nested dictionary for output formatting in the Detailed Results tab"""

    reordered_dictionary = {
        feeder: {section: {} for section in reference_dictionary[feeder]} for feeder in reference_dictionary}
    for feeder, sections in reference_dictionary.items():
        for section, stations in sections.items():
            for station in stations:
                try:
                    reordered_dictionary[feeder][section][station] = dictionary[feeder][section][station]
                except KeyError:
                    reordered_dictionary[feeder][section][station] = ''

    return reordered_dictionary


def switch_feeders(app, radial_list, select_fdrs):
    """Switch the feeder closed if it's in the feeder list, otherwise switch it open"""

    def _dic_key(my_dic, value_to_find):
        key_for_value = None
        for key, value in my_dic.items():
            if value == value_to_find:
                key_for_value = key
                break
        return key_for_value

    # get cbs of selected feeders
    select_cbs = []
    for feeder in select_fdrs:
        feeder_object = _dic_key(radial_list, feeder)
        feeder_cb = feeder_object.obj_id.obj_id
        if feeder_cb is not None and feeder_cb.GetClassName() == 'ElmCoup':
            select_cbs.extend([feeder_cb])
        else:
            select_cbs.extend(app.GetCalcRelevantObjects(feeder + '*.ElmCoup'))
    # Get all feeder cbs
    all_feeders = app.GetCalcRelevantObjects('*.ElmFeeder')
    all_cbs = []
    for feeder in all_feeders:
        all_cbs.extend(app.GetCalcRelevantObjects(feeder.loc_name + '*.ElmCoup'))

    for cb in all_cbs:
        if cb in select_cbs:
            cb.on_off = 1
        else:
            cb.on_off = 0


def get_floating_terminals(app, feeder_list, feeders_devices_section):
    """Outputs all floating terminal objects with their associated line objects for all devices for all feeders"""

    floating_terms = {feeder: {} for feeder in feeder_list}
    for feeder in feeder_list:
        feeder_elm = app.GetCalcRelevantObjects(feeder + ".ElmFeeder")
        floating_sections = floating_terminals.find_end_points(feeder_elm[0])
        for device, terminals in feeders_devices_section[feeder].items():
            floating_terms[feeder][device] = {}
            for line in floating_sections:
                t1, t2 = line.GetConnectedElements()
                t3 = line.GetConnectedElements(1,1,0)
                if len(t3) == 1 and t3[0] == t2 and t2 in terminals and t1 not in terminals:
                    floating_terms[feeder][device][line] = t1
                elif len(t3) == 1 and t3[0] == t1 and t1 in terminals and t2 not in terminals:
                    floating_terms[feeder][device][line] = t2
    app.PrintPlain(f"floating_terms: {floating_terms}")

    return floating_terms


def short_circuit(app, method, fault_type, config):
    """Perform a short circuit calculation"""

    # Get the short circuit command
    shtct = app.GetFromStudyCase('ComShc')

    # set short circuit calculation parameters
    shtct.iopt_prot = 0             # Consider protection devices. 0 equals none.
    shtct.iopt_mde = method         # Short circuit method = complete
    shtct.iopt_allbus = 1           # Busbars and junctions nodes
    shtct.iopt_shc = fault_type
    if fault_type == 'spgf':
        shtct.i_pspgf = 0           # a phase
    elif fault_type == '2psc':
        shtct.i_p2psc = 0           # a-b phases
    shtct.iopt_cur = config         # Network configuration. 0 - max, 1 - min
    shtct.ildfinit = 0              # Load Flow Initialisation
    if config == 0:                 # Voltage factor c
        shtct.cfac_full = 1.1
    elif config == 1:
        shtct.cfac_full = 1

    shtct.Execute()

    return


def sc_user_selection(app, line, node):
    """Executes a short circuit at the furtherest point on a line, assuming the node is floating"""

    # Get the short circuit command
    shtct = app.GetFromStudyCase('ComShc')

    shtct.iopt_allbus = 0           # Set fault location to User selection
    shtct.shcobj = line             # Set fault location

    shtct.iopt_opt = 'pro'          # Set line fault distance to relative
    shtct.iopt_dfr = 0              # Measure fault distance from terminal i

    if line.bus1.cterm == node:
        shtct.ppro = 1              # Set relative fault distance from Terminal i
    else:
        shtct.ppro = 99

    shtct.Execute()

    return


def _get_line_current(app, elmlne):

    def safe_attribute(line, attribute):
        if line.HasAttribute(attribute):
            current = line.GetAttribute(attribute)
        else:
            current = 0
        return current

    if elmlne.bus1 and elmlne.bus1 is not None:
        Ia1 = safe_attribute(elmlne, 'm:Ikss:bus1:A')
        Ib1 = safe_attribute(elmlne, 'm:Ikss:bus1:B')
        Ic1 = safe_attribute(elmlne, 'm:Ikss:bus1:C')

    if elmlne.bus2 and elmlne.bus2 is not None:
        Ia2 = safe_attribute(elmlne, 'm:Ikss:bus2:A')
        Ib2 = safe_attribute(elmlne, 'm:Ikss:bus2:B')
        Ic2 = safe_attribute(elmlne, 'm:Ikss:bus2:C')

    if elmlne.bus1 and elmlne.bus2 and elmlne.bus2 is not None and elmlne.bus2 is not None:
        return round(max(Ia1, Ib1, Ic1, Ia2, Ib2, Ic2), 3)
    elif elmlne.bus1 and elmlne.bus1 is not None:
        return round(max(Ia1, Ib1, Ic1), 3)
    elif elmlne.bus2 and elmlne.bus2 is not None:
        return round(max(Ia2, Ib2, Ic2), 3)
    else:
        return 0


def inrush(feeders_sections_load_total):
    """Get inrush values"""

    feeders_devices_inrush = {feeder: {} for feeder in feeders_sections_load_total}
    for feeder, devices in feeders_sections_load_total.items():
        for device, load in devices.items():
            feeders_devices_inrush[feeder][device] = \
                round((load * 12 * 1000) / (11000 * math.sqrt(3)))

    return feeders_devices_inrush,


def target_values(feeders_devices_sections, fault_type):
    """For a dictionary of nodes in a dictionary of devices in a dictionary of feeders, append the fault level
    results to the nodes"""

    if fault_type == 1:                             # P-G fault
        attribute = 'm:Ikss:A'
    elif fault_type == 2:                           # P-P fault
        attribute = 'm:Ikss:B'
    else:                                           # P-P-P fault
        attribute = 'm:Ikss'

    results_all = {feeder: {device: {} for device in feeders_devices_sections[feeder]} for feeder in
                   feeders_devices_sections}

    for feeder, devices in feeders_devices_sections.items():
        for device, nodes in devices.items():
            for node in nodes:
                results_all[feeder][device][node] = {}
                if node.HasAttribute(attribute):
                    results_all[feeder][device][node] = round(node.GetAttribute(attribute), 3)
                else:
                    results_all[feeder][device][node] = 0

    return results_all


def append_floating_terms(app, results_all, floating_terms):
    """"""

    for feeder, devices in floating_terms.items():
        for device, lines in devices.items():
            for line, terminal in lines.items():
                if line.IsEnergized():
                    sc_user_selection(app, line, terminal)
                    line_current = _get_line_current(app, line)
                else:
                    line_current = 0
                results_all[feeder][device].update({terminal: line_current})

    # Sort the dictionary in descending order
    results_sorted = {feeder: {device: {} for device in results_all[feeder]} for feeder in results_all}
    results_sorted_object = {feeder: {device: {} for device in results_all[feeder]} for feeder in results_all}
    for feeder, devices in results_all.items():
        for device, lines in devices.items():
            results_sorted_object[feeder][device] = \
                    dict(sorted(lines.items(), key=lambda item: item[1], reverse=True))
            nodes = results_sorted_object[feeder][device]
            for node in nodes:
                results_sorted[feeder][device][node.loc_name] = results_sorted_object[feeder][device][node]

    return results_sorted, results_sorted_object


def result_bounds(results_all, target):
    """Find the line node with the max/min fault level, as defined by input:
    If target = 0, min values are found. If target = 1, max values are found"""

    results = {feeder: {device: {} for device in results_all[feeder]} for feeder in results_all}
    results_locname = {feeder: {device: {} for device in results_all[feeder]} for feeder in results_all}
    for feeder, devices in results_all.items():
        for device, nodes in devices.items():
            non_zero_nodes = {x: y for x, y in nodes.items() if y != 0}
            for node in nodes:
                if target == 0:
                    try:
                        min_node = min(non_zero_nodes, key=non_zero_nodes.get)
                    except:
                        min_node = "no devices"
                    if node == min_node:
                        results[feeder][device][node] = results_all[feeder][device][node]
                        results_locname[feeder][device][node.loc_name] = results_all[feeder][device][node]
                elif target == 1:
                    try:
                        max_node = max(nodes, key=nodes.get)
                    except:
                        max_node = "no devices"
                    if node == max_node:
                        results[feeder][device][node] = results_all[feeder][device][node]
                        results_locname[feeder][device][node.loc_name] = results_all[feeder][device][node]

    return results_locname, results


def reset_min_source_imp(new_grid_data, sys_norm_min=False):
    """"""

    for grid, attributes in new_grid_data.items():
        if sys_norm_min and attributes[10] <= 0:
            grid.SetAttribute('outserv', 1)
        elif sys_norm_min:
            grid.SetAttribute('ikssmin', attributes[10])
            grid.SetAttribute('rntxnmin', attributes[11])
            grid.SetAttribute('z2tz1min', attributes[12])
            grid.SetAttribute('x0tx1min', attributes[13])
            grid.SetAttribute('r0tx0min', attributes[14])
        else:
            grid.SetAttribute('outserv', 0)
            grid.SetAttribute('ikssmin', attributes[5])
            grid.SetAttribute('rntxnmin', attributes[6])
            grid.SetAttribute('z2tz1min', attributes[7])
            grid.SetAttribute('x0tx1min', attributes[8])
            grid.SetAttribute('r0tx0min', attributes[9])
    return


def largest_tr(app, feeders_sections_loading, results_max_tr_fault):
    """Get the size of the largest transformer on the feeder section"""

    # Convert the terminal into an trfstat object
    results_max_tr_object = {
        feeder: {device: {} for device in results_max_tr_fault[feeder]} for feeder in results_max_tr_fault}
    results_max_tr = {
        feeder: {device: {} for device in results_max_tr_fault[feeder]} for feeder in results_max_tr_fault}
    for feeder, devices in results_max_tr_fault.items():
        for device, tfmrs in devices.items():
            for tfmr in tfmrs:
                trfstat = tfmr.cpSubstat
                results_max_tr_object[feeder][device][trfstat] = tfmrs.get(tfmr)
                try:
                    results_max_tr[feeder][device][trfstat.loc_name] = tfmrs.get(tfmr)
                except AttributeError:
                    results_max_tr[feeder][device][tfmr.loc_name] = tfmrs.get(tfmr)

    # Convert largest transformer furtherest upstream into an trfstat object
    feeders_sections_trmax_size = {feeder: {section: {} for section in feeders_sections_loading[feeder]} for feeder in
                            feeders_sections_loading}
    for feeder, sections in feeders_sections_loading.items():
        for section, trs in sections.items():
            tr_objects = results_max_tr_object[feeder][section]
            for key in tr_objects:
                try:
                    term_contents = key.GetContents('*.ElmTerm', 1)
                except AttributeError:
                    term_contents = []
                for tr, load in trs.items():
                    if tr in term_contents:
                        feeders_sections_trmax_size[feeder][section][key.loc_name] = load

    # A rare exception. When the trfstat object doesn't have any ElmTerm elements in its folder contents
    for feeder, sections in results_max_tr_fault.items():
        for section, transformers in sections.items():
            if not feeders_sections_trmax_size[feeder][section]:  # If the dictionary is empty
                for tfmr in transformers:
                    tfmr_objects = app.GetCalcRelevantObjects("*.ElmTerm")
                    for tr_object in tfmr_objects:
                        if tfmr in [tr_object.loc_name, tr_object]:
                            for feed, sects in feeders_sections_loading.items():
                                for sect, trs in sects.items():
                                    for tr, load in trs.items():
                                        if tr == tr_object:
                                            try:
                                                chop = tfmr[:-5]
                                            except TypeError:
                                                chop = tfmr.loc_name[:-5]
                                            feeders_sections_trmax_size[feed][sect][chop] = load

    return results_max_tr, feeders_sections_trmax_size


def output_results(app, sub_name, external_grid, feeders_devices_inrush, results_max_3p, results_max_2p,
            results_max_pg, results_min_2p, results_min_3p, results_min_pg, result_sys_norm_min_2p,
            result_sys_norm_min_pg, feeders_sections_trmax_size, results_max_tr_3p, results_max_tr_pg,
            results_all_max_3p, results_all_max_2p, results_all_max_pg, results_all_min_3p, results_all_min_2p,
            results_all_min_pg, result_all_sys_norm_min_2p, result_all_sys_norm_min_pg, feeders_devices_load,
            results_lines_max_3p, results_lines_max_2p, results_lines_max_pg, results_lines_min_3p,
            results_lines_min_2p, results_lines_min_pg, result_lines_sys_norm_min_2p, result_lines_sys_norm_min_pg,
            result_lines_type, result_lines_therm_rating, fdrs_open_switches, fdrs_lines_oos):
    """Format results file for MS Excel"""

    import openpyxl
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter
    import math
    import datetime

    app.PrintPlain("Creating output file...")

    wb = openpyxl.Workbook()

    # Set up Inputs sheet
    wb.create_sheet(index=0, title="Inputs Summary")
    wb.remove(wb["Sheet"])
    i_s = wb["Inputs Summary"]
    i_s.column_dimensions['A'].width = 19.86
    i_s.column_dimensions['B'].width = 9.14
    i_s.column_dimensions['C'].width = 9.14

    # set fonts
    heading_font = Font(size=14, bold=True)
    subheading_font = Font(size=11, bold=True)
    i_s['A1'].font = heading_font
    i_s['A3'].font = subheading_font
    i_s['A7'].font = heading_font
    i_s['A13'].font = subheading_font
    i_s['A22'].font = subheading_font

    # Set headings
    i_s['A1'] = sub_name + " Feeder Fault Level Study"
    i_s['A3'] = "Script Run Date:"
    i_s['A4'] = datetime.datetime.now()
    i_s['A7'] = "Input Summary"
    i_s['A9'] = "Short-circuit calculation method: Complete"
    i_s['A10'] = "c-Factor (max): 1.1"
    i_s['C10'] = "Voltage factor c (max): 1.1"
    i_s['A11'] = "c-Factor (min): 1.0"
    i_s['C11'] = "Voltage factor c (min): 1.0"
    i_s['A13'] = "External Grid Data:"
    i_s['A16'] = "3-P fault level (A):"
    i_s['A17'] = "R/X:"
    i_s['A18'] = "Z2/Z1:"
    i_s['A19'] = "X0/X1:"
    i_s['A20'] = "R0/X0:"
    i_s['A22'] = ("NOTE: Open points to adjacent bulk supply substations may not be "
                   "detected unless the relevant grids are active under the current "
                  "project")

    # Write external grid data
    next_column = 2
    for grid, value in external_grid.items():
        i_s.cell(column=next_column, row=14, value=grid).font = subheading_font
        i_s.cell(column=next_column, row=15, value="Maximum")
        i_s.cell(column=next_column + 1, row=15, value="Minimum")
        i_s.cell(column=next_column + 2, row=15, value="System Normal Minimum")
        next_row = 16
        for values in value[:5]:
            i_s.cell(column=next_column, row=next_row, value=values)
            next_row += 1
        next_row = 16
        for values in value[5:10]:
            i_s.cell(column=next_column + 1, row=next_row, value=values)
            next_row += 1
        next_row = 16
        for values in value[10:]:
            i_s.cell(column=next_column + 2, row=next_row, value=values)
            next_row += 1
        next_column += 4

    # Set up Results sheet
    wb.create_sheet(index=1, title="Results Summary")
    r_s = wb["Results Summary"]
    length = len(feeders_devices_inrush)
    j = 1
    for i in range(length):
        k = get_column_letter(j)
        r_s.column_dimensions[k].width = 37
        j += 5

    # Set fonts
    r_s.column_dimensions['A'].width = 36.86
    r_s['A1'].font = heading_font
    r_s['A3'].font = heading_font
    r_s['A5'].font = heading_font

    # Set headings
    r_s['A1'] = sub_name + " Feeder Fault Level Study"
    r_s['A3'] = "Results Summary"

    # Set up Detailed Results sheet
    # Set headings and fonts
    for key, value in feeders_devices_inrush.items():
        wb.create_sheet(title=key)
        feed = wb[key]
        feed['A1'] = sub_name + " Feeder Fault Level Study - Detailed Results for " + key
        feed['A1'].font = heading_font

    def nested_dic(a, b, c):
        """Template for writing nested dictionary data to the Results Summary sheet """

        next_column = 2
        for key, value in a.items():
            r_s.cell(column=next_column - 1, row=5, value="FEEDER:").font = heading_font
            r_s.cell(column=next_column, row=5, value=key).font = heading_font
            next_row = 7
            for key2, values2 in value.items():
                r_s.cell(column=next_column - 1, row=next_row, value="SECTION:").font = heading_font
                r_s.cell(column=next_column, row=next_row, value=key2).font = heading_font
                r_s.cell(column=next_column - 1, row=next_row + b, value=c)
                r_s.cell(column=next_column, row=next_row + b, value=values2)
                next_row += 14
            next_column += 5

    def two_nested_dic(a, b, c):
        """Template for writing level 2 nested dictionary data to the Results Summary sheet """

        next_column = 2
        for key, value in a.items():
            next_row = 7
            for key2, value2 in value.items():
                for key3, value3, in value2.items():
                    if key3[-5:] == "_Term":
                        key4 = key3[:-5]
                    else:
                        key4 = key3
                    r_s.cell(column=next_column + 1, row=next_row + b, value=key4)
                    r_s.cell(column=next_column - 1, row=next_row + b, value=c)
                    r_s.cell(column=next_column, row=next_row + b, value=value3)
                next_row += 14
            next_column += 5
        return

    def elements_open(fdrs_open_switches, fdrs_lines_oos, feeders_devices_inrush):
        """
        """

        feeder_row_dic = {feeder: len(feeders_devices_inrush[feeder]) for feeder in feeders_devices_inrush}

        next_column = 2
        for feeder, open_switches in fdrs_open_switches.items():
            last_row = 7 + (feeder_row_dic[feeder] * 14)
            r_s.cell(column=next_column-1, row=last_row, value="FEEDER OPEN POINTS:").font = heading_font
            if len(open_switches) > 0:
                for site, switch in open_switches.items():
                    switch_name = switch.GetAttribute("loc_name")
                    if switch.GetClassName() == "StaSwitch":
                        r_s.cell(column=next_column, row=last_row, value=switch_name)
                    else:
                        site_name = site.GetAttribute("loc_name")
                        r_s.cell(column=next_column, row=last_row, value=f"{site_name} / {switch_name}")
                    last_row += 1
            else:
                r_s.cell(column=next_column - 1, row=last_row, value="(None detected)")
                last_row += 1
            last_row += 1
            r_s.cell(column=next_column-1, row=last_row, value="FEEDER LINES OUT OF SERVICE:").font = heading_font
            feeder_lines_oos = fdrs_lines_oos[feeder]
            if len(feeder_lines_oos) > 0:
                for line_oos in feeder_lines_oos:
                    r_s.cell(column=next_column, row=last_row, value=line_oos.GetAttribute("loc_name"))
                    last_row +=1
            else:
                r_s.cell(column=next_column, row=last_row, value="(None detected)")
            next_column += 5

    def sheets_nested_dic(wb, node_data, line_data, col_node_data, col_line_data, col_name_1, col_name_2):
        """Template for writing level 2 nested dictionary data over multiple sheets
        """

        for key, value in node_data.items():
            feed = wb[key]
            feed.column_dimensions['B'].width = 18.14
            feed.column_dimensions['M'].width = 18.14
            feed.column_dimensions['W'].width = 18.14
            feed.column_dimensions['AG'].width = 18.14
            feed.column_dimensions['AQ'].width = 18.14
            next_column = 3
            for key2, value2 in value.items():
                feed.cell(column=next_column, row=3, value=key2).font = heading_font
                feed.cell(column=next_column - 2, row=3, value="SECTION:").font = heading_font
                feed.cell(column=next_column - 1, row=4, value="Downstream terminal")
                feed.cell(column=next_column + col_node_data, row=4, value=col_name_1)
                next_row = 3
                for key3, value3 in value2.items():
                    if key3[-5:] == "_Term":
                        key4 = key3[:-5]
                    else:
                        key4 = key3
                    feed.cell(column=next_column - 1, row=next_row + 2, value=key4)
                    feed.cell(column=next_column + col_node_data, row=next_row + 2, value=value3)
                    next_row += 1
                feed.cell(column=next_column - 2, row=next_row + 3, value="LINES").font = heading_font
                feed.cell(column=next_column - 1, row=next_row + 4, value="Line")
                feed.cell(column=next_column + col_line_data, row=next_row + 4, value=col_name_2)
                length = 0
                line_dict = line_data[key][key2]
                for key5, value5 in line_dict.items():
                    feed.cell(column=next_column - 1, row=next_row + length + 5, value=key5.loc_name)
                    feed.cell(column=next_column + col_line_data, row=next_row + length + 5, value=value5)
                    length += 1
                next_column += 11
        return

    # Write results to the Results sheet and the Detailed Results sheets in MS Excel
    nested_dic(feeders_devices_inrush, 1, "Inrush (A):")

    two_nested_dic(results_max_3p, 2, "Max 3-P fault level (kA) (Site):")
    two_nested_dic(results_max_2p, 3, "Max 2-P fault level (kA) (Site):")
    two_nested_dic(results_max_pg, 4, "Max P-G fault level (kA) (Site):")
    two_nested_dic(results_min_3p, 5, "Min 3-P fault level (kA) (Site):")
    two_nested_dic(results_min_2p, 6, "Min 2-P fault level (kA) (Site):")
    two_nested_dic(results_min_pg, 7, "Min P-G fault level (kA) (Site):")
    two_nested_dic(result_sys_norm_min_2p, 8, "System normal Min 2-P fault level (kA) (Site):")
    two_nested_dic(result_sys_norm_min_pg, 9, "System normal Min P-G fault level (kA) (Site):")
    two_nested_dic(feeders_sections_trmax_size, 10, "Largest transformer size (kVA) (Site):")
    two_nested_dic(results_max_tr_3p, 11, "Max 3-P at largest transformer (kA) (Site):")
    two_nested_dic(results_max_tr_pg, 12, "Max P-G at largest transformer (kA) (Site):")

    elements_open(fdrs_open_switches, fdrs_lines_oos, feeders_devices_inrush)

    sheets_nested_dic(wb, feeders_devices_load, result_lines_type, -2, -2, "Tfmr load (kVA)", "Conductor type")
    sheets_nested_dic(wb, results_all_max_3p, result_lines_therm_rating, 0, 0, "Max 3P fault", "Rated 1s current (kA)")
    sheets_nested_dic(wb, results_all_max_2p, results_lines_max_3p, 1, 1, "Max 2P fault", "Max 3P fault")
    sheets_nested_dic(wb, results_all_max_pg, results_lines_max_2p, 2, 2, "Max PG fault", "Max 2P fault")
    sheets_nested_dic(wb, results_all_min_3p, results_lines_max_pg, 3, 3, "Min 3P fault", "Max PG fault")
    sheets_nested_dic(wb, results_all_min_2p, results_lines_min_3p, 4, 4, "Min 2P fault", "Min 3P fault")
    sheets_nested_dic(wb, results_all_min_pg, results_lines_min_2p, 5, 5, "Min PG fault", "Min 2P fault")
    sheets_nested_dic(wb, result_all_sys_norm_min_2p, results_lines_min_pg, 6, 6, "Min Sys Norm 2P fault",
                      "Min PG fault")
    sheets_nested_dic(wb, result_all_sys_norm_min_pg, result_lines_sys_norm_min_2p, 7, 7, "Min Sys Norm PG fault",
                      "Min Sys Norm 2P fault")
    sheets_nested_dic(wb, results_all_max_3p, result_lines_sys_norm_min_pg, 0, 8, "Max 3P fault",
                      "Min Sys Norm PG fault")

    return wb


def save_results(app, sub_name, wb):
    """Save file to disk"""

    import os
    from pathlib import Path
    import time

    date_string = time.strftime("%Y%m%d-%H%M%S")
    file_name = sub_name + " fault level study " + date_string + ".xlsx"

    home_path = str(Path.home())
    save_path = r'\\client\Y$\PROTECTION\PowerFactory\Fault Level Studies'
    save_path_2 = '\\\\client\\' + home_path[0] + '$' + home_path[2:]
    save_path_3 = home_path

    # Try to save the output file to the SEQ Protection department drive
    try:
        wb.save(os.path.join(save_path, file_name))
        app.PrintPlain("Output file saved to " + save_path)
    # If this fails, save the output file to the user's home directory
    except:
        try:
            wb.save(os.path.join(save_path_2, file_name))
            app.PrintPlain("Output file saved to " + save_path_2)
        except:
            # When running a local install of PowerFactory
            wb.save(os.path.join(save_path_3, file_name))
            app.PrintPlain("Output file saved to " + save_path_3)

    return
